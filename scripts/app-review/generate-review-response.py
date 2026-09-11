from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUTPUT = ROOT / "docs" / "app-review" / "NeuroTrader_App_Review_Response.docx"
LOGO = ROOT / "mobile" / "assets" / "neurotrader-logo-web.png"

TEAL = "0F9F91"
DARK = "0F172A"
MUTED = "475569"
LIGHT = "E2E8F0"
PALE = "ECFDF5"


def set_cell_fill(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def add_bullet(doc, text, level=0):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    paragraph = doc.add_paragraph(style=style)
    paragraph.add_run(text)
    paragraph.paragraph_format.space_after = Pt(3)
    return paragraph


def add_step(doc, number, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.18)
    paragraph.paragraph_format.first_line_indent = Inches(-0.18)
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run(f"{number}. ")
    run.bold = True
    paragraph.add_run(text)
    return paragraph


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.55)
section.bottom_margin = Inches(0.55)
section.left_margin = Inches(0.7)
section.right_margin = Inches(0.7)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(9.5)
styles["Normal"].font.color.rgb = RGBColor.from_string(DARK)
styles["Normal"].paragraph_format.space_after = Pt(6)
styles["Title"].font.name = "Aptos Display"
styles["Title"].font.size = Pt(25)
styles["Title"].font.bold = True
styles["Title"].font.color.rgb = RGBColor.from_string(DARK)
styles["Heading 1"].font.name = "Aptos Display"
styles["Heading 1"].font.size = Pt(15)
styles["Heading 1"].font.bold = True
styles["Heading 1"].font.color.rgb = RGBColor.from_string(TEAL)
styles["Heading 1"].paragraph_format.space_before = Pt(12)
styles["Heading 1"].paragraph_format.space_after = Pt(5)

header = section.header
header_table = header.add_table(rows=1, cols=2, width=Inches(7.1))
header_table.autofit = False
header_table.columns[0].width = Inches(4.2)
header_table.columns[1].width = Inches(2.9)
header_table.cell(0, 0).paragraphs[0].add_run().add_picture(str(LOGO), width=Inches(2.85))
right = header_table.cell(0, 1).paragraphs[0]
right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
right.add_run("SG PAX Corp\n").bold = True
right.add_run("NeuroTrader Product Team\n")
right.add_run("September 10, 2026")
for cell in header_table.rows[0].cells:
    set_cell_margins(cell, 0, 0, 0, 0)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer_run = footer.add_run("SG PAX Corp  |  NeuroTrader  |  support@neurotrader-journal.com")
footer_run.font.size = Pt(8)
footer_run.font.color.rgb = RGBColor.from_string(MUTED)

title = doc.add_paragraph(style="Title")
title.add_run("App Review Response and Access Guide")
subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(12)
subtitle_run = subtitle.add_run("Submitted by SG PAX Corp for Neuro Trader iOS version 1.0.0 build 5")
subtitle_run.bold = True
subtitle_run.font.color.rgb = RGBColor.from_string(MUTED)

intro = doc.add_paragraph()
intro.add_run("To Apple App Review: ").bold = True
intro.add_run(
    "This document answers the six requested review items and provides the exact access path for the submitted iOS app. "
    "The iOS app supports sign-in to an existing NeuroTrader membership but does not offer account registration, pricing, checkout, subscription purchase, upgrade, or an external purchase link."
)

summary = doc.add_table(rows=4, cols=2)
summary.alignment = WD_TABLE_ALIGNMENT.CENTER
summary.autofit = False
summary.columns[0].width = Inches(2.05)
summary.columns[1].width = Inches(5.0)
summary_data = [
    ("App purpose", "Private trading-business planning, journaling, risk controls, analytics, simulation, and educational AI reflection"),
    ("Target audience", "Self-directed traders operating their activity as a measurable business"),
    ("Commerce in iOS", "None. Memberships are managed outside the submitted app"),
    ("Trading activity", "No order routing, trade execution, custody, brokerage, or financial advice"),
]
for row, values in zip(summary.rows, summary_data):
    for index, value in enumerate(values):
        cell = row.cells[index]
        cell.text = value
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        if index == 0:
            set_cell_fill(cell, PALE)
            cell.paragraphs[0].runs[0].bold = True

heading = doc.add_heading("1 Physical device screen recording", level=1)
keep_with_next(heading)
doc.add_paragraph(
    "A separate uninterrupted recording captured on a physical iPhone running the latest available iOS version is attached to the Resolution Center reply. It begins from the iPhone Home Screen with a cold app launch and shows the normal review flow."
)
add_bullet(doc, "Launch Neuro Trader and show that the first screen is Sign in, with no account-creation or payment control.")
add_bullet(doc, "Log in with the main review account and open Business Center, Trading Business Plan, Execution Journal, P&L, KPIs, Business Notebook, AI Coach, Settings, legal links, and sign out.")
add_bullet(doc, "Show paid-feature access through the preauthorized review account. No purchase is performed in the iOS app.")
add_bullet(doc, "Log in with the disposable deletion account and complete Settings, Danger zone, Delete account, including the email, DELETE phrase, and both confirmations.")
add_bullet(doc, "The app has no public feed or public user-generated content. Journals, notebook entries, screenshots, and AI conversations are private to the authenticated account, so public reporting and blocking mechanisms do not apply.")

heading = doc.add_heading("2 Purpose target audience problem and value", level=1)
keep_with_next(heading)
doc.add_paragraph(
    "NeuroTrader is a trading-business operating and educational tool for self-directed Trader Entrepreneurs. It addresses fragmented plans, spreadsheets, journal notes, risk rules, and performance reviews by connecting them in one private operating workspace. Users can define a Trading Business Plan, record executions and decisions, monitor personal risk controls, review measurable performance, organize operating knowledge, simulate future paths, and request objective AI-assisted reflection grounded in their own records."
)
doc.add_paragraph(
    "NeuroTrader does not execute trades, route orders, hold funds or securities, operate a brokerage, recommend the purchase, sale, or holding of a security, or guarantee income, profit, capital growth, or any projected result. Analytics, simulations, projections, and AI outputs are educational and may differ materially from actual results."
)

heading = doc.add_heading("3 Setup access and main feature instructions", level=1)
keep_with_next(heading)
add_step(doc, 1, "Install and launch the latest submitted build. No sample-file upload is required.")
add_step(doc, 2, "At Sign in, use the main review email appreview@neurotrader-journal.com and the password entered in App Store Connect Review Notes.")
add_step(doc, 3, "Business Center opens with the preloaded Two-Year Growth Demo personal account.")
add_step(doc, 4, "Open Trading Business Plan to review the simulated $10,000 starting balance, $250,000 target, two-year dates, risk limits, checkpoints, and completed-cycle continuation recommendation.")
add_step(doc, 5, "Open P&L and September 7, 2026 in Execution Journal to inspect populated stock, listed option-contract, crypto, and forex transactions.")
add_step(doc, 6, "Open KPIs, Business Notebook, and Coach. A suggested coach question is Analyze my two-year plan objectively.")
add_step(doc, 7, "Open Settings to inspect notification consent, Privacy Policy, Terms and Conditions, support, sign out, workspace reset, and permanent account deletion.")
add_step(doc, 8, "For deletion testing, sign out and use appreview-delete@neurotrader-journal.com with the same review password. Complete the Delete account flow in Settings.")

credentials = doc.add_table(rows=4, cols=2)
credentials.alignment = WD_TABLE_ALIGNMENT.CENTER
credentials.style = "Table Grid"
credential_data = [
    ("Main review email", "appreview@neurotrader-journal.com"),
    ("Deletion test email", "appreview-delete@neurotrader-journal.com"),
    ("Password", "Enter the Keychain-stored review password in App Store Connect only"),
    ("Sample files", "Not required; synthetic review data is preloaded"),
]
for row, values in zip(credentials.rows, credential_data):
    for index, value in enumerate(values):
        row.cells[index].text = value
        set_cell_margins(row.cells[index])
        if index == 0:
            set_cell_fill(row.cells[index], PALE)
            row.cells[index].paragraphs[0].runs[0].bold = True

heading = doc.add_heading("4 External services tools and platforms", level=1)
keep_with_next(heading)
services = [
    "Supabase provides authentication, access status, the private application database, row-level access control, and private file storage.",
    "Vercel hosts the authenticated application API used by the mobile client and scheduled operational jobs.",
    "OpenAI API generates user-requested educational AI Coach responses from the selected private records and plan context disclosed in the app.",
    "Expo Notifications and Apple Push Notification service provide optional device registration, business reminders, and separately opted-in promotional notifications.",
    "Resend sends account, security, support, and service-related email outside the iOS app.",
    "Stripe manages memberships initiated outside the submitted iOS app. The iOS app contains no pricing, checkout, purchase, upgrade, or external purchase-link flow.",
    "SnapTrade and Webull integration code supports planned direct broker connectivity, but direct broker connections are disabled in the submitted iOS build pending provider approvals. The app does not execute or route trades.",
]
for service in services:
    add_bullet(doc, service)

heading = doc.add_heading("5 Regional differences", level=1)
keep_with_next(heading)
doc.add_paragraph(
    "The submitted app provides the same core features and content in all supported regions. Users may choose English or Spanish. Optional notifications depend on the user's device permission and Apple service availability. Direct broker connections are disabled in this build, so there is no regional broker-feature difference in the submitted binary."
)

heading = doc.add_heading("6 Regulated activity and protected third party material", level=1)
keep_with_next(heading)
doc.add_paragraph(
    "NeuroTrader is educational recordkeeping, simulation, analytics, and business-discipline software. It is not a broker-dealer, investment adviser, exchange, custodian, or trade-execution service. It does not provide individualized financial advice or promise any trading or capital outcome."
)
doc.add_paragraph(
    "The submitted iOS app does not display, reproduce, or distribute third-party curriculum books, paid market publications, or other protected training material. Product copy, workflows, calculations, and synthetic sample records shown in the app are owned by SG PAX Corp or were created specifically for the product. No market-data redistribution or protected third-party content is included in the submitted mobile experience."
)

heading = doc.add_heading("Reviewer support and final submission notes", level=1)
keep_with_next(heading)
add_bullet(doc, "Support email: support@neurotrader-journal.com")
add_bullet(doc, "Support URL: https://www.neurotrader-journal.com/contact")
add_bullet(doc, "Privacy Policy URL: https://www.neurotrader-journal.com/privacy")
add_bullet(doc, "Terms and Conditions URL: https://www.neurotrader-journal.com/terms")
add_bullet(doc, "All preloaded balances, trades, journals, metrics, and outcomes are synthetic and labeled for demonstration and App Review only.")

closing = doc.add_paragraph()
closing.paragraph_format.space_before = Pt(10)
closing.add_run("Submitted by SG PAX Corp\n").bold = True
closing.add_run("NeuroTrader Product Team")

doc.core_properties.title = "NeuroTrader App Review Response and Access Guide"
doc.core_properties.subject = "Apple App Review requested information"
doc.core_properties.author = "SG PAX Corp"
doc.core_properties.keywords = "NeuroTrader, App Review, iOS, access guide"
doc.save(OUTPUT)
print(OUTPUT)
